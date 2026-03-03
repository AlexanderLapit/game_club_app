from docx import Document
from datetime import datetime

def create_test_case_report(test_cases):
    """
    Создаёт файл 'ТестКейс.docx' с результатами тестов.
    """
    doc = Document()
    doc.add_heading('Отчёт по тестированию системы', 0)
    doc.add_paragraph(f"Автоматически сгенерировано: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Тест-кейс'
    hdr_cells[1].text = 'Результат'
    hdr_cells[2].text = 'Время'

    for case in test_cases:
        row = table.add_row().cells
        row[0].text = case['case']
        row[1].text = case['result']
        row[2].text = case['time']

    doc.save('ТестКейс.docx')
    print("Отчёт 'ТестКейс.docx' успешно создан.")